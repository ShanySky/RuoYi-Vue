"""通过实际 Pi、若依、数据库与隔离服务完成本人数据到三类私有成果的闭环。"""

from io import BytesIO
import hashlib
import json
import threading
import time
import uuid
from zipfile import ZipFile
import requests
import ai_server_api_acceptance as h
from ai_server_governance_acceptance import sql, steps
from ai_evolution_stack import credentials

PROGRAM = """python - <<'PY'
import json
from openpyxl import Workbook
from docx import Document
from reportlab.pdfgen.canvas import Canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
rows=json.load(open('input.json'))['rows']
wb=Workbook();sheet=wb.active;sheet.title='本人范围'
sheet.append(['用户编号','用户名','部门'])
document=Document();document.add_heading('本人授权范围用户',0)
document.add_paragraph('人数：'+str(len(rows)))
table=document.add_table(rows=1,cols=3)
for cell,text in zip(table.rows[0].cells,['用户编号','用户名','部门']):cell.text=text
pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
pdf=Canvas('users.pdf',pageCompression=0);pdf.setFont('STSong-Light',16);pdf.drawString(50,790,'本人授权范围用户')
pdf.setFont('Helvetica',10);pdf.drawString(50,765,'Count: '+str(len(rows)))
for i,row in enumerate(rows):
 values=[row['user_id'],row['user_name'],row['dept_name']]
 sheet.append(values)
 for cell,text in zip(table.add_row().cells,values):cell.text=str(text)
 pdf.drawString(50,740-i*18,str(row['user_id'])+' '+row['user_name'])
wb.save('users.xlsx');document.save('users.docx');pdf.save()
print(json.dumps({'count':len(rows),'files':['users.xlsx','users.docx','users.pdf']}))
PY"""


def run():
    h.REQUEST_TIMEOUT = 360
    root = h.api(None, "/login", "POST", {"username": "admin", "password": "admin123"})["token"]
    tag = "work_" + uuid.uuid4().hex[:7]
    users, roles, stored = [], [], []
    evidence = {"passed": False, "checks": [], "trace": str(h.TRACE), "paidRequests": 0}
    auth = {"Authorization": "Bearer " + credentials()["workspace"]}
    folder = h.EVIDENCE / ("workspace-fullstack-" + h.STAMP)
    folder.mkdir()

    def broker(action, body):
        response = requests.post("http://127.0.0.1:28097/v1/" + action, headers=auth, json=body, timeout=45)
        assert response.status_code == 200, response.text
        return response.json()

    def identity(suffix, dept, permissions):
        name = tag + suffix
        menus = sql("select menu_id from sys_menu where perms in (" + ",".join(["%s"] * len(permissions)) + ")", permissions)
        h.api(root, "/system/role", "POST", {"roleName": name, "roleKey": name, "roleSort": 90, "dataScope": "3", "status": "0",
            "menuIds": [row["menu_id"] for row in menus]})
        role = sql("select role_id from sys_role where role_key=%s", (name,))[0]["role_id"]
        roles.append(role)
        h.api(root, "/system/user", "POST", {"userName": name, "nickName": name, "deptId": dept, "password": "TestOnly_123!", "roleIds": [role], "postIds": [], "status": "0"})
        user = sql("select user_id from sys_user where user_name=%s", (name,))[0]["user_id"]
        users.append(user)
        token = h.api(None, "/login", "POST", {"username": name, "password": "TestOnly_123!"})["token"]
        return token, role, user, name

    def policy(enabled):
        current = h.api(root, "/ai/admin/workspace")["data"]
        return h.api(root, "/ai/admin/workspace", "PUT", {"enabled": enabled, "revision": current["revision"]})

    def note(name, **values):
        evidence["checks"].append({"name": name, **values})
        print(json.dumps(evidence["checks"][-1], ensure_ascii=False), flush=True)

    try:
        h.api(root, "/ai/config/provider", "POST", {"name": "框架工作空间确定性验收", "baseUrl": f"http://127.0.0.1:{h.MODEL_PORT}/v1", "token": "isolated-mock-only", "enabled": True, "timeoutSeconds": 60})
        model = next(row for row in h.api(root, "/ai/config/models")["data"] if row["modelCode"] == "mock-api-model")
        request = {"modelId": model["modelId"], "route": "/index", "pageContext": {}, "frontendTools": []}
        reader, role, uid, name = identity("reader", 103, ("system:user:list", "ai:workspace:use"))
        other, _, _, _ = identity("outside", 104, ("system:user:list", "ai:workspace:use"))
        manager, _, _, _ = identity("manager", 103, ("ai:workspace:view", "ai:workspace:edit"))
        policy(False)
        h.api(reader, "/ai/chat/turn", "POST", {**request, "userMessage": tag + "closed"})
        assert not any(item["function"]["name"].startswith("workspace_") for item in h.REQUESTS[-1]["tools"])
        policy(True)
        h.api(manager, "/ai/chat/turn", "POST", {**request, "userMessage": tag + "manager"})
        assert not any(item["function"]["name"].startswith("workspace_") for item in h.REQUESTS[-1]["tools"])
        assert h.api(manager, "/ai/admin/workspace")["data"]["environment"]["ready"]
        note("默认关闭与治理权不等于执行权")

        native = h.api(reader, "/system/user/list?userName=" + tag)["rows"]
        assert len(native) == 2 and uid in {row["userId"] for row in native}
        assert all(row["userName"] != tag + "outside" for row in native)
        sequence = [("server_data_describe", {"id": "data_users"}),
            ("data_users", {"operation": "QUERY", "columns": ["user_id", "user_name", "dept_name"], "filters": [{"field": "user_name", "operator": "contains", "value": tag}], "limit": 100}),
            ("workspace_import_result", lambda results: {"resultId": json.loads(results[1]["content"])["result"]["resultId"], "path": "input.json"}),
            ("workspace_execute", {"command": PROGRAM}),
            *[("workspace_publish", {"path": "users." + extension, "name": "本人范围用户." + extension}) for extension in ("xlsx", "docx", "pdf")]]
        h.SCENARIOS[tag + "combined"] = steps(sequence)
        response = h.api(reader, "/ai/chat/turn", "POST", {**request, "userMessage": tag + "combined"})["data"]
        assert response["type"] == "MESSAGE", response
        conversation = response["conversationId"]
        result_calls = sql("select tool_name,status,tool_result_json from ai_server_call where conversation_id=%s order by create_time,call_id", (conversation,))
        assert len(result_calls) == len(sequence) and all(row["status"] == "SUCCEEDED" for row in result_calls), result_calls
        files = h.api(reader, f"/ai/chat/conversations/{conversation}/artifacts")["data"]
        assert len(files) == 3
        for artifact in files:
            stored.append({"userId": uid, "runId": artifact["runId"], "id": artifact["id"]})
            downloaded = requests.get(h.BASE + f"/ai/artifacts/{artifact['id']}/download", headers={"Authorization": "Bearer " + reader}, timeout=65)
            assert downloaded.headers.get("Cache-Control") == "no-store"
            assert downloaded.headers["Content-Disposition"].startswith("attachment;")
            data = downloaded.content
            assert len(data) == artifact["bytes"] and hashlib.sha256(data).hexdigest() == artifact["sha256"]
            extension = artifact["name"].rsplit(".", 1)[1]
            (folder / ("users." + extension)).write_bytes(data)
            if extension in ("xlsx", "docx"):
                with ZipFile(BytesIO(data)) as package:
                    body = package.read("xl/worksheets/sheet1.xml" if extension == "xlsx" else "word/document.xml").decode()
                    assert all(row["userName"] in body for row in native) and tag + "outside" not in body
            else:
                assert data.startswith(b"%PDF-") and data.rstrip().endswith(b"%%EOF") and all(row["userName"].encode() in data for row in native)
                assert ("Count: " + str(len(native))).encode() in data
            denied = h.api(other, f"/ai/artifacts/{artifact['id']}/download", expect=None)
            assert denied["code"] != 200
        note("本人数据经实际 Pi 导入隔离空间，三类文件内容、哈希和授权下载通过", conversation=conversation,
            artifacts=[{key: file[key] for key in ("id", "name", "bytes", "sha256")} for file in files])
        assert h.api(other, f"/ai/chat/conversations/{conversation}/artifacts", expect=None)["code"] != 200
        # 不改变旧令牌，直接撤销原数据权限；成果和历史必须同时关闭。
        sql("delete rm from sys_role_menu rm join sys_menu m on m.menu_id=rm.menu_id where rm.role_id=%s and m.perms='system:user:list'", (role,))
        for artifact in files:
            assert h.api(reader, f"/ai/artifacts/{artifact['id']}/download", expect=None)["code"] != 200
        assert h.api(reader, f"/ai/chat/conversations/{conversation}", expect=None)["code"] != 200
        before = len(h.REQUESTS)
        assert h.api(reader, "/ai/chat/turn", "POST", {**request, "conversationId": conversation, "userMessage": "继续读取旧成果"}, expect=None)["code"] != 200
        assert len(h.REQUESTS) == before
        note("撤销原业务权限后旧令牌不能下载、读取历史或把旧内容发给模型")
        deadline = time.monotonic() + 12
        while broker("status", {})["tasks"] and time.monotonic() < deadline:
            time.sleep(1)
        assert broker("status", {})["tasks"] == 0
        evidence["passed"] = True
        note("任务结束后临时容器回收，已发布成果独立保留")
    finally:
        for user in users:
            for row in sql("select artifact_id,run_id from ai_artifact where user_id=%s", (user,)):
                item = {"userId": user, "runId": row["run_id"], "id": row["artifact_id"]}
                if item not in stored:
                    stored.append(item)
        for artifact in stored:
            broker("delete", artifact)
            sql("delete from ai_artifact where artifact_id=%s", (artifact["id"],))
        for user in users:
            h.api(root, "/system/user/" + str(user), "DELETE")
        for role in roles:
            h.api(root, "/system/role/" + str(role), "DELETE")
        evidence["fixtureCleanup"] = True
        (folder / "evidence.json").write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding="utf-8")
    print("AI_WORKSPACE_FULLSTACK_OK " + str(folder), flush=True)


if __name__ == "__main__":
    server = h.ThreadingHTTPServer(("127.0.0.1", h.MODEL_PORT), h.Model)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    try:
        run()
    finally:
        server.shutdown()
        server.server_close()
